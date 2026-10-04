"""v9 (trimmed): one rule for every slide —
   text-only slides  → subheading + ▣ section + ◆ bullets (bold key phrases, inline citations)
   slides with visual → text on one side, figure/table on the other (left/right or top/bottom)
   visuals are real (PX4 User Guide photos, CC BY 4.0) or grounded (dataset values, PX4 EKF2-style chain, Gantt).
Base = user's v2 deck → 4:3, VDEC banner, title + underline. Fonts: Latin Arial / East Asian MS PGothic."""
import os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x00, 0x0F, 0x65); MID = RGBColor(0x33, 0x33, 0x99); PALE = RGBColor(0xDA, 0xED, 0xEF)
BAR = RGBColor(0x8C, 0x96, 0xC9); TEXT = RGBColor(0x33, 0x33, 0x33); BLACK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66); WHITE = RGBColor(0xFF, 0xFF, 0xFF); LINEC = RGBColor(0xBF, 0xBF, 0xBF)
SOFT = RGBColor(0xC8, 0xCE, 0xEA)
EA = "MS PGothic"
L, R, W = 0.5, 9.5, 9.0

p = Presentation(SRC)
S = list(p.slides)


# ------------------------------------------------------------------ primitives
def font(run, size, bold=False, color=TEXT, italic=False):
    f = run.font
    f.name = "Arial"; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for e in rPr.findall(qn("a:ea")): rPr.remove(e)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)


def parts(t):
    return [(x, i % 2 == 1) for i, x in enumerate(re.split(r"\*\*", t)) if x]


def bullet(pg, char, marL=0.0, indent=0.0, pct=80):
    pPr = pg._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(marL)))); pPr.set("indent", str(int(Inches(-indent))))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buSzPct"):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if char is None:
        etree.SubElement(pPr, qn("a:buNone")); return
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(pct * 1000))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", EA)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def nostyle(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None: sh._element.remove(st)


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    nostyle(sh)
    return sh


def fill_text(sh, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, m=(0.08, 0.04)):
    """paras: [(text, size, color, bold)]"""
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(m[0]); tf.margin_top = tf.margin_bottom = Inches(m[1])
    for k, (t, sz, c, b) in enumerate(paras):
        pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pg.alignment = align; pg.line_spacing = 1.1; bullet(pg, None)
        r = pg.add_run(); r.text = t; font(r, sz, b, c)
    return sh


def tbox(slide, x, y, w, h, paras, **kw):
    return fill_text(slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)), paras, **kw)


def arrow(slide, x1, y1, x2, y2, color=MID, width=1.75):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def body(slide, x, y, w, h, items, base=15):
    """items: (kind, text) — sub | sec | b | b2 | cite ; **bold** markup."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"): setattr(tf, "margin_" + side, Inches(0.02))
    for k, (kind, t) in enumerate(items):
        pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pg.line_spacing = 1.15
        if kind == "sub":
            bullet(pg, None); pg.space_after = Pt(4)
            for x_, b in parts(t): r = pg.add_run(); r.text = x_; font(r, base + 4, True, BLACK)
        elif kind == "sec":
            bullet(pg, "▣", 0.3, 0.3, 90); pg.space_before = Pt(0 if k == 0 else 10); pg.space_after = Pt(2)
            for x_, b in parts(t): r = pg.add_run(); r.text = x_; font(r, base + 1, True, NAVY)
        elif kind == "b":
            bullet(pg, "◆", 0.62, 0.24, 75); pg.space_before = Pt(4)
            for x_, b in parts(t): r = pg.add_run(); r.text = x_; font(r, base, b, BLACK if b else TEXT)
        elif kind == "b2":
            bullet(pg, "–", 0.9, 0.18, 100); pg.space_before = Pt(1)
            for x_, b in parts(t): r = pg.add_run(); r.text = x_; font(r, base - 0.5, b, BLACK if b else TEXT)
        elif kind == "cite":   # citation on its own line under a bullet
            bullet(pg, None, 0.62, 0); pg.space_before = Pt(0)
            r = pg.add_run(); r.text = t; font(r, base - 4, False, MUTED)
    return tb


def subhead(slide, t):
    body(slide, L, 1.1, W, 0.5, [("sub", t)])


def table(slide, x, y, w, colw, rows, size=12.5, rowh=0.38):
    gt = slide.shapes.add_table(len(rows), len(colw), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    tbl = gt.table; tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None: tblPr.remove(sid)
    for c, wv in zip(tbl.columns, colw): c.width = Inches(wv)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j); hdr = i == 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            pg = cell.text_frame.paragraphs[0]; pg.alignment = PP_ALIGN.CENTER
            r = pg.add_run(); r.text = val
            font(r, size, hdr or j == 0, NAVY if hdr else (BLACK if j == 0 else TEXT))
            tcPr = cell._tc.get_or_add_tcPr()
            for side, wpt, col in (("lnL", 0, None), ("lnR", 0, None),
                                   ("lnT", 1.5 if hdr else 0, "000F65" if hdr else None),
                                   ("lnB", 1.0 if hdr or i == len(rows) - 1 else 0.5, "000F65" if hdr or i == len(rows) - 1 else "BFBFBF")):
                ln = etree.SubElement(tcPr, qn("a:" + side)); ln.set("w", str(int(Pt(wpt))))
                if wpt == 0: etree.SubElement(ln, qn("a:noFill"))
                else:
                    sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", col)
            if hdr:
                sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", "E8ECF5")
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    return gt


def clear(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder and "TITLE" in str(sh.placeholder_format.type): continue
        sh._element.getparent().remove(sh._element)


def title(slide, t):
    tf = slide.shapes.title.text_frame
    tf.paragraphs[0].runs[0].text = t
    for r in tf.paragraphs[0].runs[1:]: r.text = ""
    for r in tf.paragraphs[0].runs:
        rPr = r._r.get_or_add_rPr()
        for e in rPr.findall(qn("a:ea")): rPr.remove(e)
        etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def photo(slide, f, x, y, w):
    from PIL import Image
    im = Image.open(f); h = w * im.size[1] / im.size[0]
    slide.shapes.add_picture(f, Inches(x), Inches(y), Inches(w), Inches(h))
    return h


# ================================================================== 1 표지
s = S[0]; clear(s)
t = s.shapes.title; t.text_frame.text = ""
fill_text(t, [("GPS 사용 불가 환경에서 비전 보조 항법을 위한", 26, NAVY, True), ("다중센서 융합", 26, NAVY, True)])
t.left, t.top, t.width, t.height = Inches(0.6), Inches(1.15), Inches(8.8), Inches(1.45)
tbox(s, 0.8, 2.65, 8.4, 0.45, [("측정 주기와 지연을 고려한 카메라·IMU·거리센서 융합 추정기", 17, MID, True)])
sh = tbox(s, 0.8, 3.15, 8.4, 0.5, [("Multi-Sensor Fusion for Vision-Aided Navigation in GPS-Denied Environments", 13, MID, False)])
sh.text_frame.paragraphs[0].runs[0].font.italic = True
tbox(s, 1.0, 4.35, 8.0, 1.75, [("학생자율연구 II  연구계획 발표  |  2026학년도 2학기", 14, NAVY, True),
                               ("박찬혁 (202121212)  ·  지도교수: 안창선 교수님", 14, TEXT, False),
                               ("School of Mechanical Engineering", 13, TEXT, False), ("Pusan National University", 13, TEXT, False)])
notes(s, "[10초]\n안녕하세요, 학생자율연구II 연구계획을 발표할 박찬혁입니다. 제 연구는 GPS를 쓰기 어려운 환경에서 카메라, IMU, 거리센서를 하나의 필터로 융합하되, 센서마다 다른 측정 시점까지 고려하는 위치·자세 추정기를 만드는 것입니다.")

# ================================================================== 2 연구 배경  (text left · one real photo right)
s = S[2]; clear(s); title(s, "연구 배경")
subhead(s, "GPS가 끊기는 환경에서는 카메라로 항법을 보조한다")
body(s, L, 1.75, 5.4, 4.9, [
    ("sec", "GPS 제한 환경의 항법"),
    ("b", "실내·터널에서는 위성 신호가 끊겨 차량·드론 모두 **IMU(관성측정장치)**에 의존함"),
    ("b", "IMU는 가속도를 적분하므로 **오차가 시간에 따라 누적**됨"),
    ("cite", "(Groves, 2013)"),
    ("sec", "비전 보조 항법"),
    ("b", "**카메라 영상의 움직임**을 IMU와 결합하여 위치·자세를 추정함"),
    ("cite", "(Mourikis and Roumeliotis, 2007)"),
    ("b", "드론은 **하향 카메라 + 거리센서**로 이동을 측정함 (예: PX4, NASA Ingenuity)"),
    ("cite", "(PX4 User Guide; Bayard et al., 2019)"),
], base=15)
px = 6.15; pw = R - px
h1 = photo(s, f"{HERE}/photos/flow_lidar_attached.jpg", px, 2.35, pw)
tbox(s, px, 2.35 + h1 + 0.05, pw, 0.5, [("드론 아래에 장착한 하향 카메라(PX4Flow)와 거리센서(Lidar-Lite)", 10.5, TEXT, True)], align=PP_ALIGN.LEFT, m=(0, 0))
tbox(s, px, 2.35 + h1 + 0.55, pw, 0.25, [("사진: PX4 User Guide (CC BY 4.0)", 9, MUTED, False)], align=PP_ALIGN.LEFT, m=(0, 0))
notes(s, "[50초]\n자동차 내비게이션은 터널에서 GPS가 끊겨도 위치를 계속 보여 줍니다. IMU, 즉 관성 정보로 위치를 이어서 계산하기 때문인데, 가속도를 적분하는 방식이라 시간이 지날수록 오차가 쌓입니다. 실내를 나는 드론도 같은 문제를 겪습니다.\n"
         "그래서 카메라로 움직임을 관측해 IMU 오차를 잡아 주는 비전 보조 항법을 씁니다. (오른쪽 사진) 드론은 이렇게 아래를 보는 카메라와 거리센서를 달아 이동을 측정하고, 드론 비행제어기 PX4와 NASA의 화성 헬리콥터 인저뉴어티도 같은 조합을 씁니다. 그런데 이 센서들을 실제로 합치려면 풀어야 할 문제가 있습니다.\n\n"
         "사진 출처: PX4 User Guide, Optical Flow 문서 (CC BY 4.0)")

# ================================================================== 3 필요성 및 목적 (text left · one table right · purpose below)
s = S[3]; clear(s); title(s, "연구의 필요성 및 목적")
subhead(s, "세 센서를 융합해야 하지만, 측정 시점이 서로 다르다")
body(s, L, 1.75, 5.35, 4.0, [
    ("sec", "세 센서를 함께 쓰는 이유"),
    ("b", "**IMU**: 시간이 지날수록 오차가 누적됨"),
    ("b", "**카메라**: 실제 이동 거리를 모름 → **거리센서가 필요**"),
    ("cite", "(Groves, 2013; PX4 User Guide)"),
    ("sec", "문제: 측정 시점의 차이"),
    ("b", "센서마다 **측정 주기가 다르고**, 영상은 처리 시간만큼 **늦게 도착**함"),
    ("b", "이 시점 차이는 **융합 성능을 크게 저하**함"),
    ("cite", "(Qin and Shen, 2018)"),
], base=15)
tbox(s, 6.1, 1.85, 3.4, 0.32, [("공개 드론 데이터의 센서 주기", 12.5, NAVY, True)], align=PP_ALIGN.LEFT, m=(0, 0))
table(s, 6.1, 2.22, 3.4, [1.45, 0.9, 1.05], [
    ["센서", "주기", "측정 간격"],
    ["IMU", "200 Hz", "5 ms"],
    ["하향 카메라", "20 Hz", "50 ms"],
    ["레이저 거리계", "30 Hz", "약 33 ms"],
], size=12, rowh=0.44)
tbox(s, 6.1, 4.05, 3.4, 0.3, [("INSANE 데이터셋 (Brommer et al., 2024)", 9.5, MUTED, False)], align=PP_ALIGN.LEFT, m=(0, 0))
g = box(s, L, 5.45, W, 0.8, fill=None, line=NAVY, lw=1.5)
tf = g.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.2)
pg = tf.paragraphs[0]; pg.alignment = PP_ALIGN.LEFT
for t_, sz, b, c in [("연구 목적   ", 14.5, True, NAVY), ("측정 주기와 지연 차이를 고려하여 카메라·IMU·거리센서를 ", 14, False, TEXT),
                     ("하나의 필터로 융합하는 위치·자세 추정기", 14, True, BLACK), ("를 구현", 14, False, TEXT)]:
    r = pg.add_run(); r.text = t_; font(r, sz, b, c)
notes(s, "[50초]\n세 센서를 함께 쓰는 이유는 각자 한계가 있기 때문입니다. IMU는 오차가 쌓이고, 카메라는 영상 속 이동만 알려 줘서 실제로 몇 미터 움직였는지는 거리센서가 있어야 압니다.\n"
         "문제는 측정 시점입니다. (오른쪽 표) 제가 사용할 공개 드론 데이터에서 IMU는 5 ms마다, 카메라는 50 ms마다, 거리센서는 약 33 ms마다 측정됩니다. 여기에 영상 처리 시간까지 더해지면, 지금 들어온 영상 측정은 사실 조금 전의 움직임입니다. 이런 시점 차이는 융합 성능을 크게 떨어뜨린다고 보고되어 있고, 실제 드론 비행제어기 PX4도 센서별 지연값을 따로 두고 보정합니다.\n"
         "그래서 이 연구의 목적은 측정 주기와 지연 차이를 고려해 세 센서를 하나의 필터로 융합하는 위치·자세 추정기를 구현하는 것입니다.")

# ================================================================== 4 선행연구 (text only)
s = S[4]; clear(s); title(s, "선행연구와 적용 방향")
subhead(s, "차량 분야의 시간 처리 방법을 비전 항법에 적용")
body(s, L, 1.75, W, 4.9, [
    ("sec", "레이더·V2X 기반 선행 차량 상태 추정 (Kim and Ahn, 2025)"),
    ("b", "도착 시점이 서로 다른 레이더와 V2X 정보를 융합하여 앞 차량의 상태를 추정함"),
    ("b", "본 연구에서 참고할 **시간 처리 방법**"),
    ("b2", "**지연 보상**: 데이터의 수신 시점과 실제 측정 시점을 구분"),
    ("b2", "**시간 정렬**: 함께 사용하는 신호의 시점을 맞추고, 필요하면 보간"),
    ("b2", "**과거 상태 재추정**: 늦게 온 측정으로 과거 값을 보정한 후 현재까지 다시 추정"),
    ("sec", "드론 항법의 지연 측정 처리"),
    ("b", "MSF-EKF는 지연 측정이 오면 **과거 상태에서 다시 예측**하여 반영함 (Lynen et al., 2013)"),
    ("sec", "본 연구의 적용 방향"),
    ("b", "차량 모델이 아닌 **시간 처리 방법**을 카메라·IMU·거리센서 구성에 적용하고, **적용 전·후를 비교**함"),
], base=15)
notes(s, "[45초]\n이 시간 문제는 차량 분야에서도 다뤄졌습니다. 지도교수님 연구실의 Kim과 Ahn의 논문은 레이더와 V2X 통신처럼 도착 시점이 다른 정보를 융합해 앞 차량의 상태를 추정했습니다. 여기서 지연 보상, 시간 정렬, 과거 상태 재추정 세 가지를 참고합니다.\n"
         "드론 항법에서도 MSF-EKF처럼, 지연 측정이 오면 과거 상태로 돌아가 다시 예측하는 구조가 쓰입니다. 저는 차량 모델이 아니라 이 시간 처리 방법을 제 센서 구성에 옮겨서, 적용 전과 후를 비교하겠습니다.")

# ================================================================== 5 연구방법 (text top · figure bottom)
s = S[5]; clear(s); title(s, "연구방법")
subhead(s, "EKF 기반 다중센서 융합과 지연 측정 처리")
body(s, L, 1.75, W, 2.6, [
    ("sec", "기본 융합"),
    ("b", "영상 이동은 **KLT**(특징점 추적, Lucas and Kanade, 1981)로 구하고, **EKF(확장 칼만필터)**로 융합함"),
    ("b", "**예측**: IMU  /  **보정**: 영상 이동(회전 성분 제거)과 거리"),
    ("sec", "시간 처리"),
    ("b", "측정값을 **버퍼에 저장**하고, 늦게 온 측정은 **측정 시점에서 보정 후 현재까지 재추정**함"),
], base=15)
fy = 4.45
for k, (nm, sub_) in enumerate([("IMU", "가속도·각속도"), ("하향 카메라", "KLT 영상 이동"), ("거리센서", "지면까지 거리")]):
    y = fy + k * 0.62
    fill_text(box(s, L, y, 1.85, 0.52, fill=WHITE, line=MID), [(nm, 12.5, NAVY, True), (sub_, 10, TEXT, False)])
    arrow(s, L + 1.85, y + 0.26, 2.75, fy + 0.83)
fill_text(box(s, 2.75, fy + 0.18, 1.95, 1.3, fill=PALE), [("버퍼", 14, NAVY, True), ("측정 시점 기준", 11, TEXT, False), ("저장·정렬", 11, TEXT, False)])
arrow(s, 4.7, fy + 0.83, 5.15, fy + 0.83)
fill_text(box(s, 5.15, fy + 0.03, 2.2, 1.6, fill=NAVY), [("EKF", 18, WHITE, True), ("예측 · 보정", 12, WHITE, False), ("늦은 측정 → 재추정", 11, SOFT, False)])
arrow(s, 7.35, fy + 0.83, 7.8, fy + 0.83)
fill_text(box(s, 7.8, fy + 0.33, 1.7, 1.0, fill=WHITE, line=MID), [("위치·자세", 14, NAVY, True)])
tbox(s, L, 6.4, W, 0.28, [("구조 참고: PX4 EKF2, MSF-EKF (Lynen et al., 2013)", 9.5, MUTED, False)], align=PP_ALIGN.LEFT, m=(0, 0))
notes(s, "[50초]\n방법입니다. 영상 이동은 KLT 특징점 추적으로 구하고, 확장 칼만필터, 즉 EKF로 융합합니다. IMU로 예측하고, 영상 이동과 거리로 보정하는데, 영상 이동에 섞인 회전 성분은 자이로로 빼고 거리로 실제 크기를 정합니다.\n"
         "(아래 그림) 시간 처리를 위해 센서 데이터를 측정 시점 기준으로 버퍼에 저장해 둡니다. 늦게 도착한 측정은 그 측정 시점으로 돌아가 보정하고, 저장해 둔 IMU 입력으로 현재까지 다시 계산합니다. 이 구조는 PX4 비행제어기와 MSF-EKF에서도 쓰는 방식입니다.")

# ================================================================== 6 구현 및 평가 (text only)
s = p.slides.add_slide(p.slide_layouts[1])
for ph in list(s.placeholders):
    if "TITLE" not in str(ph.placeholder_format.type): ph._element.getparent().remove(ph._element)
s.shapes.title.text_frame.text = "구현 및 평가 계획"
for r in s.shapes.title.text_frame.paragraphs[0].runs:
    rPr = r._r.get_or_add_rPr(); etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)
subhead(s, "공개 드론 데이터로 시간 처리의 효과를 검증")
body(s, L, 1.75, W, 4.9, [
    ("sec", "데이터"),
    ("b", "**INSANE**: 하향 카메라·레이저 거리계·IMU, 실내 구간의 **모션캡처 정답 궤적** (Brommer et al., 2024)"),
    ("b", "늦게 도착하는 상황 재현: **도착 시점 = 촬영 시점 + 실측 처리시간 + 추가 지연**"),
    ("sec", "구현 도구"),
    ("b", "**Python·OpenCV**: 영상 이동(KLT)과 처리시간 계산"),
    ("b", "**MATLAB**: EKF·버퍼·재추정 구현 및 평가"),
    ("sec", "비교 및 평가"),
    ("b", "**A. 지연 무시** vs **B. 시간 처리**를 추가 지연 0·25·50·100 ms에서 비교"),
    ("b", "지표: **ATE**(위치 RMSE), **RE**(구간별 상대 오차), **처리시간** (Zhang and Scaramuzza, 2018)"),
], base=15)
notes(s, "[45초]\n구현과 평가 계획입니다. 데이터는 공개 드론 데이터셋 INSANE을 씁니다. 하향 카메라, 레이저 거리계, IMU가 모두 있고, 실내 구간에는 모션캡처로 측정한 정답 궤적이 있습니다. 이 데이터의 카메라 시간은 실제 촬영 시점이라, 늦게 도착하는 상황은 촬영 시점에 실측한 처리시간과 추가 지연을 더해 만듭니다.\n"
         "영상 처리는 파이썬 OpenCV로, 필터와 시간 처리는 매트랩으로 구현합니다. 지연을 무시한 경우와 시간 처리를 한 경우를 추가 지연 0에서 100 ms까지 비교하고, 정답 궤적 대비 위치 오차와 처리시간으로 평가하겠습니다.")
lst = p.slides._sldIdLst
ids = lst.findall(qn("p:sldId")); new = ids[-1]; lst.remove(new); lst.insert(6, new)

# ================================================================== 7 연구계획 (Gantt)
s = S[6]
old = next(sh for sh in s.shapes if sh.has_table)
plan = [[old.table.cell(i, j).text_frame.text.replace("\x0b", "\n") for j in range(3)] for i in range(1, 5)]
clear(s); title(s, "연구계획")
subhead(s, "기본 융합 → 시간 처리 → 비교 검증 (10월 ~ 12월)")
gx0, ncol = 5.0, 6
gcw = (R - gx0) / ncol
tbase = 1.85
for j, m in enumerate(["10월", "11월", "12월"]):
    tbox(s, gx0 + 2 * j * gcw, tbase, 2 * gcw, 0.26, [(m, 12.5, NAVY, True)], m=(0, 0))
for j in range(ncol):
    tbox(s, gx0 + j * gcw, tbase + 0.27, gcw, 0.22, [("전반" if j % 2 == 0 else "후반", 10, MUTED, False)], m=(0, 0))
rh, ry0 = 1.0, tbase + 0.55
spans = [(0, 1), (1, 2), (2, 4), (4, 6)]
for i, (when, what, res) in enumerate(plan):
    y = ry0 + i * rh
    if i % 2 == 0: box(s, L, y, W, rh, fill=PALE)
    wl = [l_ for l_ in what.split("\n") if l_.strip()]
    tbox(s, L + 0.1, y, gx0 - L - 0.15, rh, [(wl[0], 14.5, NAVY, True), (" · ".join(wl[1:]), 12, TEXT, False)],
         align=PP_ALIGN.LEFT, m=(0, 0))
for j in range(ncol + 1):
    ln = s.shapes.add_connector(1, Inches(gx0 + j * gcw), Inches(tbase + 0.27), Inches(gx0 + j * gcw), Inches(ry0 + 4 * rh))
    ln.line.color.rgb = LINEC; ln.line.width = Pt(0.75 if j % 2 == 0 else 0.5)
    if j % 2: ln.line.dash_style = 4
for i, (c0, c1) in enumerate(spans):
    y = ry0 + i * rh
    box(s, gx0 + c0 * gcw + 0.04, y + rh / 2 - 0.15, (c1 - c0) * gcw - 0.08, 0.3, fill=NAVY if i == 2 else BAR)
notes(s, "[25초]\n일정입니다. 10월 전반에는 선행논문과 센서 데이터를 분석하고, 10월 후반에 옵티컬 플로우와 IMU 예측으로 기본 융합 구조를 만듭니다. 핵심인 11월에는 지연 보상과 시간 정렬, 과거 상태 재추정을 넣고, 12월에 적용 전후를 비교해 보고서로 정리하겠습니다.")

# ================================================================== 8 참고문헌 (numbered text)
s = S[7]; clear(s); title(s, "참고문헌")
refs = [
    "M. Kim and C. Ahn, “State Estimation of Preceding Target Vehicle Using Radar and V2X,” IEEE Access, vol. 13, pp. 198482–198495, 2025.",
    "P. D. Groves, Principles of GNSS, Inertial, and Multisensor Integrated Navigation Systems, 2nd ed., Artech House, 2013.",
    "A. I. Mourikis and S. I. Roumeliotis, “A Multi-State Constraint Kalman Filter for Vision-aided Inertial Navigation,” IEEE ICRA, 2007.",
    "D. S. Bayard et al., “Vision-Based Navigation for the NASA Mars Helicopter,” AIAA SciTech Forum, 2019.",
    "T. Qin and S. Shen, “Online Temporal Calibration for Monocular Visual-Inertial Systems,” IEEE/RSJ IROS, 2018.",
    "S. Lynen et al., “A Robust and Modular Multi-Sensor Fusion Approach Applied to MAV Navigation,” IEEE/RSJ IROS, 2013.",
    "B. D. Lucas and T. Kanade, “An Iterative Image Registration Technique with an Application to Stereo Vision,” IJCAI, 1981.",
    "C. Brommer et al., “The INSANE Dataset: Large Number of Sensors for Challenging UAV Flights in Mars Analog, Outdoor, and Out-/Indoor Transition Scenarios,” IJRR, vol. 43, no. 8, 2024.",
    "Z. Zhang and D. Scaramuzza, “A Tutorial on Quantitative Trajectory Evaluation for Visual(-Inertial) Odometry,” IEEE/RSJ IROS, 2018.",
    "PX4 Autopilot User Guide, “Optical Flow” and “Using PX4's Navigation Filter (EKF2)”, docs.px4.io (CC BY 4.0), 2026. 10. 확인.",
]
tb = s.shapes.add_textbox(Inches(L), Inches(1.2), Inches(W), Inches(4.9))
tf = tb.text_frame; tf.word_wrap = True
for k, rt in enumerate(refs):
    pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    bullet(pg, None, 0.42, 0.42); pg.space_before = Pt(0 if k == 0 else 6); pg.line_spacing = 1.05
    r = pg.add_run(); r.text = f"[{k + 1}]  "; font(r, 11.5, True, NAVY)
    r = pg.add_run(); r.text = rt; font(r, 11.5, False, TEXT)
tbox(s, L, 6.15, W, 0.5, [("감사합니다", 20, NAVY, True)])
notes(s, "[15초]\n정리하면, 차량 분야에서 다뤄 온 시간 처리 방법을 카메라·IMU·거리센서 기반 비전 항법에 적용하고 그 효과를 확인하는 것이 이 연구입니다. 이상으로 발표를 마치겠습니다. 의견과 질문 부탁드립니다. 감사합니다.")

# ------------------------------------------------------------------ delete 목차
lst = p.slides._sldIdLst
sid = lst.findall(qn("p:sldId"))[1]
p.part.drop_rel(sid.get(qn("r:id"))); lst.remove(sid)
p.save(OUT)
print("saved", OUT)
