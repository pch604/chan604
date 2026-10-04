"""Apply the v3 수정안 to the user's 8-slide v2 deck (4:3), keeping its master/banner and style."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
NAVY = RGBColor(0x00, 0x00, 0x66)
TEXT = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0xF1, 0xF1, 0xF1)
MID = RGBColor(0x80, 0x80, 0x80)
TINT = RGBColor(0xE8, 0xEA, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
EA = "맑은 고딕"

p = Presentation(SRC)
S = list(p.slides)


def set_font(run, size=None, bold=None, color=None):
    f = run.font
    f.name = "Arial"
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for e in rPr.findall(qn("a:ea")): rPr.remove(e)
    ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", EA)


def tb(slide, x, y, w, h, paras, size=15, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       shape=None, fill=None, line=None, name=None, margin=0.05):
    """paras: list of str or list of (text, opts) runs per paragraph."""
    if shape is None:
        sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill is None: sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        if line is None: sh.line.fill.background()
        else:
            sh.line.color.rgb = line; sh.line.width = Pt(1.25)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.12
    if name: sh.name = name
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pg.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = pg.add_run(); r.text = t
            set_font(r, o.get("size", size), o.get("bold", bold), o.get("color", color))
    return sh


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=1.5, dash=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def remove(shape):
    shape._element.getparent().remove(shape._element)


def by_name(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_cell(cell, paras, size=15, bold_first=True, color=TEXT):
    tf = cell.text_frame
    # keep first paragraph's pPr, clear the rest
    first = tf.paragraphs[0]
    for pg in list(tf.paragraphs)[1:]:
        pg._p.getparent().remove(pg._p)
    for r in list(first.runs): r._r.getparent().remove(r._r)
    for br in first._p.findall(qn("a:br")): first._p.remove(br)
    if isinstance(paras, str): paras = [paras]
    for i, t in enumerate(paras):
        if i == 0:
            pg = first
        else:
            new = copy.deepcopy(first._p)
            for r in new.findall(qn("a:r")): new.remove(r)
            for e in new.findall(qn("a:endParaRPr")): new.remove(e)
            first._p.getparent().append(new)
            pg = tf.paragraphs[-1]
        r = pg.add_run(); r.text = t
        set_font(r, size, bold_first and i == 0, NAVY if (bold_first and i == 0) else color)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- slide 1 표지
s = S[0]
tb(s, 0.9, 3.38, 8.2, 0.5, "— 측정 주기와 지연을 고려한 카메라·IMU·거리센서 융합 추정기", size=18, color=NAVY,
   align=PP_ALIGN.CENTER, name="부제")
info = next(sh for sh in s.shapes if sh.name == "Rectangle 4")
paras = info.text_frame.paragraphs
texts = [pg.text for pg in paras]
order = [1, 0, 2]  # 학생자율연구II 연구계획 발표 → 박찬혁 → 지도교수
new_texts = [texts[i] for i in order]
for pg, t in zip(paras, new_texts):
    runs = pg.runs
    runs[0].text = t
    for r in runs[1:]: r.text = ""
notes(s, "[10초]\n안녕하세요, 학생자율연구II 연구계획을 발표할 박찬혁입니다. 제 연구는 GPS를 쓰기 어려운 환경에서 카메라, IMU, 거리센서를 하나의 필터로 융합하되, 센서마다 다른 측정 시점까지 고려하는 위치·자세 추정기를 만드는 것입니다.")

# ---------------------------------------------------------------- slide 3 연구 배경
s = S[2]
for n in ["항법의 공통 과제", "비전 기반 항법", "기법 소개", "편집 가능한 비교표", "필요성으로 연결"]:
    remove(by_name(s, n))
tb(s, 0.5, 1.12, 9.0, 0.5, [[("항법의 기본 구조   ", {"color": TEXT, "bold": False}), ("IMU로 예측  →  외부 센서로 보정", {})]],
   size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
   shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name="기본 구조")
cards = [("차량", ["터널에서 GPS가 끊겨도", "관성·영상 정보로", "위치를 이어서 추정"]),
         ("드론", ["GPS 없는 실내 비행", "하향 카메라 · IMU ·", "거리센서로 위치 유지"]),
         ("항공기", ["화성 헬리콥터 Ingenuity", "하향 카메라 · IMU ·", "레이저 거리계로 항법"])]
cw, gap, cy = 2.7, 0.45, 1.85
for i, (head, body) in enumerate(cards):
    x = 0.5 + i * (cw + gap)
    tb(s, x, cy, cw, 0.45, head, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
       shape=MSO_SHAPE.RECTANGLE, fill=NAVY, name=f"플랫폼 {head}")
    tb(s, x, cy + 0.45, cw, 1.15, body, size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
       shape=MSO_SHAPE.RECTANGLE, fill=None, line=NAVY, name=f"플랫폼 {head} 설명")
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw + 0.08), Inches(cy + 0.75), Inches(0.29), Inches(0.4))
        a.fill.solid(); a.fill.fore_color.rgb = MID; a.line.fill.background(); a.shadow.inherit = False
tb(s, 0.5, 3.6, 9.0, 0.35, "카메라 정보의 처리 흐름", size=15, bold=True, color=NAVY, name="흐름 제목")
fl = [("카메라 영상", None, 1.55, GRAY, None),
      ("Optical Flow", "영상 속 움직임 측정", 2.35, WHITE, NAVY),
      ("VIO · 필터", "IMU와 결합해 추정", 2.35, WHITE, NAVY),
      ("위치 · 자세", None, 1.55, GRAY, None)]
x, fy, fh = 0.5, 4.0, 0.85
for i, (t, sub, w, fill, ln) in enumerate(fl):
    paras = [[(t, {"bold": True, "color": NAVY})]] + ([[(sub, {"size": 13})]] if sub else [])
    tb(s, x, fy, w, fh, paras, size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
       shape=MSO_SHAPE.RECTANGLE, fill=fill, line=ln, name=f"흐름 {t}")
    if i < 3:
        arrow(s, x + w + 0.04, fy + fh / 2, x + w + 0.36, fy + fh / 2)
    x += w + 0.4
tb(s, 0.5, 5.15, 9.0, 1.05,
   [[("공통 과제:  센서마다 측정 주기와 지연이 다르다", {"bold": True, "size": 18})],
    [("→ 시간 차이를 고려한 융합 구조가 필요", {"size": 16})]],
   color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=NAVY, name="공통 과제")
tb(s, 0.5, 6.3, 9.0, 0.3, "Ingenuity: Bayard et al., AIAA SciTech 2019 · VIO 예: VINS-Fusion, PX4 VIO 문서",
   size=10, color=MID, name="출처")
notes(s, "[65초]\n자동차 내비게이션은 터널에 들어가 GPS가 끊겨도 화면 속 차량이 계속 움직입니다. 마지막 위치에서 속도와 회전 정보로 위치를 이어서 계산하기 때문입니다. 이처럼 항법은 관성 정보로 움직임을 예측하고, 다른 센서로 그 예측을 보정하는 구조입니다.\n"
         "보정에 많이 쓰이는 센서가 카메라입니다. 옵티컬 플로우는 연속된 영상에서 점이 얼마나 움직였는지 추적하는 기법으로, 광마우스가 움직임을 읽는 것과 같은 원리입니다. VIO는 이 영상 정보를 IMU와 결합해 위치와 자세를 추정합니다. 옵티컬 플로우가 측정이라면, VIO는 그 측정을 쓰는 추정입니다.\n"
         "이 원리는 드론과 항공기에도 그대로 쓰입니다. 실내 드론은 하향 카메라와 IMU, 거리센서로 위치를 유지하고, NASA의 화성 헬리콥터 인저뉴어티도 같은 조합으로 비행했습니다. 플랫폼은 달라도 IMU로 예측하고, 카메라와 거리센서로 보정한다는 구조는 같습니다.\n\n"
         "출처: Bayard et al., Vision-Based Navigation for the NASA Mars Helicopter, AIAA SciTech 2019 / https://github.com/HKUST-Aerial-Robotics/VINS-Fusion / https://docs.px4.io/main/en/computer_vision/visual_inertial_odometry")

# ---------------------------------------------------------------- slide 4 필요성 및 목적
s = S[3]
tbl_sh = by_name(s, "편집 가능한 비교표")
tbl_sh.top = Inches(1.25)
tbl = tbl_sh.table
tbl.rows[0].height = Inches(0.42)
tbl.rows[1].height = Inches(1.15)
tbl.rows[2].height = Inches(1.15)
remove(by_name(s, "목적 표제")); remove(by_name(s, "목적"))
# timeline figure
fy0 = 4.02
tb(s, 0.5, fy0, 5.0, 0.3, "센서별 측정 시점 (개념도)", size=13, bold=True, color=NAVY, name="타임라인 제목")
lx0, lx1 = 1.45, 7.55
ry_imu, ry_cam = fy0 + 0.45, fy0 + 0.95
tb(s, 0.5, ry_imu - 0.16, 0.9, 0.32, "IMU", size=13, bold=True, color=TEXT)
tb(s, 0.5, ry_cam - 0.16, 0.9, 0.32, "카메라", size=13, bold=True, color=TEXT)
arrow(s, lx0, ry_imu, lx1, ry_imu, color=MID, width=1, head=True)
arrow(s, lx0, ry_cam, lx1, ry_cam, color=MID, width=1, head=True)
xx = lx0 + 0.1
while xx < lx1 - 0.15:
    arrow(s, xx, ry_imu - 0.09, xx, ry_imu + 0.09, color=NAVY, width=1, head=False)
    xx += 0.2
for k, xc in enumerate([1.55, 3.35, 5.15, 6.95]):
    arrow(s, xc, ry_cam - 0.14, xc, ry_cam + 0.14, color=NAVY, width=2.25, head=False)
# delay annotation on one camera sample
arrow(s, 3.35, ry_cam + 0.2, 4.2, ry_cam + 0.2, color=NAVY, width=1.25, dash=True)
tb(s, 2.8, ry_cam + 0.24, 1.1, 0.25, "측정 시점", size=10.5, color=MID, align=PP_ALIGN.CENTER)
tb(s, 3.65, ry_cam + 0.24, 1.1, 0.25, "도착 시점", size=10.5, color=MID, align=PP_ALIGN.CENTER)
tb(s, 4.27, ry_cam + 0.07, 0.75, 0.25, "지연 Δt", size=11.5, bold=True, color=NAVY)
tb(s, 7.8, fy0 + 0.35, 1.7, 1.1, [[("예) 5 m/s × 50 ms", {"bold": True, "color": NAVY})], "= 25 cm 차이",
                                    [("단순 계산 예시", {"size": 10.5, "color": MID})]],
   size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name="지연 예시")
tb(s, 0.5, 5.72, 9.0, 0.95,
   [[("연구 목적", {"bold": True, "color": NAVY, "size": 16})],
    "측정 주기와 지연 차이를 고려하여 카메라·IMU·거리센서를 하나의 필터로 융합하는 위치·자세 추정기를 구현한다."],
   size=15, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=TINT, line=NAVY, name="연구 목적", margin=0.15)
notes(s, "[55초]\n그런데 이 구조를 실제 센서에 적용하면 두 가지 문제가 생깁니다.\n"
         "첫째는 개별 센서의 한계입니다. IMU만 쓰면 작은 오차가 적분되면서 계속 쌓입니다. 카메라는 영상 속 이동만 알려 주기 때문에, 실제 이동 거리로 바꾸려면 거리센서 정보가 필요합니다. 그래서 세 센서를 함께 써야 합니다.\n"
         "둘째는 측정 시점의 차이입니다. 아래 그림처럼 IMU는 매우 빠른 주기로 들어오지만 카메라는 그보다 느리고, 영상 처리에도 시간이 걸립니다. 그래서 지금 필터에 들어온 영상 측정은 사실 조금 전의 움직임입니다. 예를 들어 초속 5 m로 움직일 때 50 ms 늦은 측정을 현재 값처럼 쓰면 25 cm의 차이가 생깁니다.\n"
         "그래서 제 연구의 목적은 측정 주기와 지연 차이를 고려해 카메라, IMU, 거리센서를 하나의 필터로 융합하는 위치·자세 추정기를 구현하는 것입니다.")

# ---------------------------------------------------------------- slide 5 선행연구
s = S[4]
s.shapes.title.text_frame.paragraphs[0].runs[0].text = "선행연구와 적용 방향"
for r in s.shapes.title.text_frame.paragraphs[0].runs[1:]: r.text = ""
remove(by_name(s, "선행연구 방향"))
tb(s, 0.5, 1.12, 3.6, 0.8, [[("차량 분야", {"bold": True, "color": NAVY})], [("레이더 · V2X 융합 (Kim & Ahn, 2025)", {"size": 13})]],
   size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name="차량 분야")
tb(s, 5.9, 1.12, 3.6, 0.8, [[("본 연구", {"bold": True, "color": WHITE})], [("카메라 · IMU · 거리센서 융합", {"size": 13, "color": WHITE})]],
   size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=NAVY, name="본 연구")
arrow(s, 4.2, 1.52, 5.8, 1.52, width=2)
tb(s, 4.1, 1.15, 1.8, 0.3, "시간 처리 방법 적용", size=11.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
tbl_sh = by_name(s, "편집 가능한 비교표")
tbl_sh.top = Inches(2.15)
tbl = tbl_sh.table
grid = tbl._tbl.tblGrid
newcol = copy.deepcopy(grid.findall(qn("a:gridCol"))[-1]); grid.append(newcol)
for tr in tbl._tbl.tr_lst:
    tr.append(copy.deepcopy(tr.tc_lst[-1]))
widths = [2.0, 4.9, 2.1]
for c, w in zip(tbl.columns, widths): c.width = Inches(w)
hdr = tbl.cell(0, 2)
for r in hdr.text_frame.paragraphs[0].runs[1:]: r.text = ""
hdr.text_frame.paragraphs[0].runs[0].text = "근거"
for i, t in enumerate(["Kim & Ahn (2025)", "Kim & Ahn (2025)", "Kim & Ahn (2025)\nMSCKF (2007)"], start=1):
    set_cell(tbl.cell(i, 2), t.split("\n"), size=13, bold_first=False)
notes(s, "[55초]\n이 시간 문제는 차량 분야에서도 중요하게 다뤄졌습니다. 지도교수님 연구실의 Kim과 Ahn의 논문은 레이더와 V2X 통신처럼 도착 시점이 서로 다른 정보를 융합해 앞 차량의 상태를 추정했습니다. 저는 이 논문에서 세 가지 처리 방법을 참고하겠습니다.\n"
         "첫째, 지연 보상입니다. 데이터를 받은 시점과 실제로 측정한 시점을 구분합니다. 둘째, 시간 정렬입니다. 함께 쓰는 신호의 시점을 맞추고, 필요하면 앞뒤 데이터로 그 사이 값을 보간합니다. 셋째, 과거 상태 재추정입니다. 늦게 들어온 측정으로 그 시점의 추정값을 고치고, 거기서부터 현재까지 다시 계산합니다.\n"
         "비전 항법 쪽에서는 과거 카메라 자세를 함께 유지하며 추정하는 MSCKF와 OpenVINS의 구조를 참고하겠습니다. 차량 모델을 그대로 가져오는 것이 아니라, 시간 처리 방법을 제 센서 구성에 맞게 옮기는 것이 핵심입니다.")

# ---------------------------------------------------------------- slide 6 연구방법
s = S[5]
m = by_name(s, "방법 개요")
m.top = Inches(1.02)
for pg in m.text_frame.paragraphs:
    for r in pg.runs:
        r.text = r.text.replace("를 기본 구조로 검토", "를 기본 구조로 사용").replace("기본 구조로 검토", "기본 구조로 사용")
t1 = [sh for sh in s.shapes if sh.has_table]
t_main, t_val = sorted(t1, key=lambda z: z.top)
t_main.top = Inches(1.78)
set_cell(t_main.table.cell(1, 1), ["영상 이동과 거리 정보로 보정", "(자이로로 회전 성분 제거)"], size=14, bold_first=False)
remove(by_name(s, "시간 처리 설명"))
# delayed-measurement timeline
fy = 3.18
tb(s, 0.5, fy, 6.8, 0.3, "늦게 도착한 측정의 처리", size=14, bold=True, color=NAVY, name="처리 그림 제목")
ax0, ax1, ay = 0.7, 7.0, fy + 1.3
arrow(s, ax0, ay, ax1, ay, color=MID, width=1.25)
tb(s, ax1 - 0.1, ay + 0.03, 0.5, 0.25, "시간", size=10.5, color=MID)
xx = ax0 + 0.2
while xx < ax1 - 0.2:
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xx - 0.045), Inches(ay - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid(); d.fill.fore_color.rgb = NAVY; d.line.fill.background(); d.shadow.inherit = False
    xx += 0.4
tb(s, ax0, ay - 0.36, 1.3, 0.25, "IMU 예측", size=10.5, color=MID)
tk, tn = 2.3, 5.1
arrow(s, tk, ay - 0.95, tk, ay + 0.1, color=NAVY, width=1, dash=True, head=False)
arrow(s, tn, ay - 0.95, tn, ay + 0.1, color=NAVY, width=1, dash=True, head=False)
tb(s, tk - 0.9, ay + 0.1, 1.8, 0.28, "t_k  카메라 측정 시점", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
tb(s, tn - 0.9, ay + 0.1, 1.8, 0.28, "t_k+Δt  측정 도착", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
arrow(s, tn - 0.05, ay - 0.78, tk + 0.05, ay - 0.78, color=NAVY, width=2)
tb(s, tk + 0.1, ay - 1.07, tn - tk - 0.2, 0.27, "① 측정 시점으로 돌아가 보정", size=11.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
arrow(s, tk + 0.05, ay - 0.3, 6.75, ay - 0.3, color=NAVY, width=1.5, dash=True)
tb(s, tn + 0.08, ay - 0.6, 1.9, 0.27, "② 현재까지 재추정", size=11.5, bold=True, color=NAVY)
tb(s, 0.5, ay + 0.42, 6.9, 0.27, "새 측정이 없는 시점에는 예측만 진행하고, 측정이 도착하면 보정", size=12, color=TEXT)
tb(s, 7.45, fy + 0.15, 2.05, 1.75,
   [[("버퍼에 저장", {"bold": True, "color": NAVY, "size": 14})], "· 과거 상태", "· IMU 입력", "· 측정값", "· 타임스탬프"],
   size=13, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name="버퍼", margin=0.15)
# validation table: 3 data rows
t_val.top = Inches(5.22)
vt = t_val.table
row_xml = vt._tbl.tr_lst[1]
for _ in range(2):
    vt._tbl.append(copy.deepcopy(row_xml))
vt.cell(0, 0).text_frame.paragraphs[0].runs[0].text = "검증 항목"
rows = [("비교", "시간 처리 미적용 vs 적용 (같은 데이터 · 초기 조건)"),
        ("지표", "기준값 대비 위치 · 자세 오차, 재추정 처리시간"),
        ("기준값", "사용 데이터의 정답 궤적 (데이터 확정 시 출처 명시)")]
for i, (a, b) in enumerate(rows, start=1):
    set_cell(vt.cell(i, 0), [a], size=12.5, bold_first=True)
    set_cell(vt.cell(i, 1), [b], size=12.5, bold_first=False)
    vt.rows[i].height = Inches(0.32)
vt.rows[0].height = Inches(0.32)
notes(s, "[65초]\n구체적인 방법입니다. 영상의 이동은 KLT 옵티컬 플로우로 추적하고, 융합은 확장 칼만필터, 즉 EKF를 기본 구조로 사용합니다.\n"
         "예측 단계에서는 IMU의 가속도와 각속도로 다음 순간의 위치와 자세를 계산합니다. 보정 단계에서는 영상 이동과 거리 정보를 측정 모델로 넣어 예측을 바로잡습니다. 이때 영상 이동에는 회전 성분이 섞여 있어서, 자이로로 회전을 빼고 거리로 실제 크기를 정합니다.\n"
         "시간 처리는 아래 그림처럼 합니다. 먼저 센서별 주기와 타임스탬프, 지연을 확인합니다. 필터는 과거 일정 구간의 상태와 입력을 저장해 두고, 새 측정이 없을 때는 예측만 합니다. 늦게 도착한 측정이 오면 그 측정 시점으로 돌아가 보정하고, 저장해 둔 IMU 입력으로 현재까지 다시 계산합니다.\n"
         "검증은 같은 데이터와 초기 조건에서 시간 처리를 넣기 전과 후를 비교합니다. 기준값 대비 위치 오차와 재추정에 드는 처리시간을 함께 보고, 효과와 한계를 정리하겠습니다.")

# ---------------------------------------------------------------- slide 7 연구계획
s = S[6]
old = by_name(s, "편집 가능한 비교표")
cells = old.table
plan = []
for i in range(1, 5):
    plan.append([cells.cell(i, j).text_frame.text.replace("\x0b", "\n") for j in range(3)])
remove(old); remove(by_name(s, "계획 범위"))
cw, step, cy = 2.4, 2.2, 1.3
for i, (when, what, res) in enumerate(plan):
    x = 0.5 + i * step
    ch = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, Inches(x), Inches(cy), Inches(cw), Inches(0.8))
    ch.shadow.inherit = False
    ch.fill.solid(); ch.fill.fore_color.rgb = NAVY if i == 2 else TINT; ch.line.fill.background()
    tf = ch.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = when.replace("\n", " "); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_font(r, 17, True, WHITE if i == 2 else NAVY)
    lines = [l for l in what.split("\n") if l.strip()]
    tb(s, x + 0.05, cy + 1.0, 2.1, 1.45, [[(lines[0], {"bold": True, "color": NAVY})]] + lines[1:],
       size=13.5, name=f"수행 {when}")
    rl = [l for l in res.split("\n") if l.strip()]
    tb(s, x + 0.05, cy + 2.6, 2.05, 1.05, [[("확인할 결과", {"bold": True, "size": 11, "color": MID})]] + rl,
       size=12.5, shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name=f"결과 {when}", margin=0.1)
tb(s, 0.5, 5.3, 9.0, 0.75, [[("최종 목표   ", {"size": 14}), ("추정기 구현과 시간 처리 효과 검증", {"size": 18})]],
   bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=NAVY, name="최종 목표")
notes(s, "[25초]\n일정입니다. 10월 전반에는 선행논문을 분석하고 센서의 주기와 타임스탬프, 지연을 정리합니다. 10월 후반에는 옵티컬 플로우와 IMU 예측으로 기본 융합 구조를 만듭니다. 11월에는 지연 보상과 시간 정렬, 과거 상태 재추정을 넣고, 12월에 적용 전후를 비교해 보고서로 정리하겠습니다.")

# ---------------------------------------------------------------- slide 8 참고문헌
s = S[7]
ref = by_name(s, "참고문헌")
tf = ref.text_frame
last = tf.paragraphs[-1]
newp = copy.deepcopy(last._p)
last._p.addnext(newp)
pg = tf.paragraphs[-1]
runs = pg.runs
for r, t in zip(runs, ["[4] D. S. Bayard et al. (2019)", "Vision-Based Navigation for the NASA Mars Helicopter.", "AIAA SciTech 2019 Forum."]):
    r.text = t
for pgi in tf.paragraphs:
    for r in pgi.runs:
        if r.font.size: r.font.size = Pt(13)
ref.top = Inches(1.15)
ref.height = Inches(3.45)
tb(s, 0.5, 4.5, 9.0, 0.5, "차량 분야의 시간 처리 방법  →  카메라·IMU·거리센서 비전 항법에 적용", size=15, bold=True, color=NAVY,
   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, fill=GRAY, name="한 줄 요약")
notes(s, "[15초]\n정리하면, 차량 분야에서 다뤄 온 시간 처리 방법을 카메라·IMU·거리센서 기반 비전 항법에 적용하고 그 효과를 확인하는 것이 이 연구입니다. 이상으로 발표를 마치겠습니다. 의견과 질문 부탁드립니다. 감사합니다.")

# ---------------------------------------------------------------- delete 목차 (slide 2)
sldIdLst = p.slides._sldIdLst
sid = sldIdLst.findall(qn("p:sldId"))[1]
rid = sid.get(qn("r:id"))
p.part.drop_rel(rid)
sldIdLst.remove(sid)

p.save(OUT)
print("saved", OUT)
