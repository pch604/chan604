"""v4: text-led slides (▣ / ◆ style) on the user's v2 deck. Keeps 4:3, VDEC banner, title + underline layout.
Fonts: Latin Arial, East Asian MS PGothic."""
import copy, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
NAVY = RGBColor(0x00, 0x00, 0x66)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
TEXT = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
EA = "MS PGothic"

p = Presentation(SRC)
S = list(p.slides)


def font(run, size, bold=False, color=TEXT):
    f = run.font
    f.name = "Arial"; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)


def runs_of(text):
    """'**bold**' markup → [(text, bold)]"""
    out = []
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part: out.append((part, i % 2 == 1))
    return out


def bullet(pg, char, marL, indent, size_pct=80):
    pPr = pg._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(marL)))); pPr.set("indent", str(int(Inches(-indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buSzPct"):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if char is None:
        etree.SubElement(pPr, qn("a:buNone")); return
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(size_pct * 1000))
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", EA)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def spacing(pg, before=0, after=0, line=1.15):
    pg.space_before = Pt(before); pg.space_after = Pt(after); pg.line_spacing = line


def body(slide, x, y, w, h, items, base=15):
    """items: ("sub", text) subheading · ("sec", text) ▣ header · ("b", text) ◆ bullet · ("b2", text) – sub-bullet · ("p", text) plain"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"): setattr(tf, "margin_" + side, Inches(0.02))
    first = True
    for kind, text in items:
        pg = tf.paragraphs[0] if first else tf.add_paragraph()
        if kind == "sub":
            bullet(pg, None, 0, 0); spacing(pg, 0 if first else 6, 4, 1.1)
            for t, b in runs_of(text): font(pg.add_run(), base + 4, True, BLACK)
            for r, (t, b) in zip(pg.runs, runs_of(text)): r.text = t
        elif kind == "sec":
            bullet(pg, "▣", 0.3, 0.3, 90); spacing(pg, 0 if first else 9, 2, 1.1)
            for t, b in runs_of(text):
                r = pg.add_run(); r.text = t; font(r, base + 1, True, NAVY)
        elif kind == "b":
            bullet(pg, "◆", 0.65, 0.25, 75); spacing(pg, 3, 0, 1.15)
            for t, b in runs_of(text):
                r = pg.add_run(); r.text = t; font(r, base, b, BLACK if b else TEXT)
        elif kind == "b2":
            bullet(pg, "–", 0.95, 0.2, 100); spacing(pg, 1, 0, 1.12)
            for t, b in runs_of(text):
                r = pg.add_run(); r.text = t; font(r, base - 0.5, b, BLACK if b else TEXT)
        else:
            bullet(pg, None, 0, 0); spacing(pg, 3, 0, 1.15)
            for t, b in runs_of(text):
                r = pg.add_run(); r.text = t; font(r, base, b, BLACK if b else TEXT)
        first = False
    return tb


def clear(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
            continue
        sh._element.getparent().remove(sh._element)


def title(slide, text):
    tf = slide.shapes.title.text_frame
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]: r.text = ""
    for r in tf.paragraphs[0].runs:
        rPr = r._r.get_or_add_rPr()
        for e in rPr.findall(qn("a:ea")): rPr.remove(e)
        etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ 1 표지
s = S[0]
info = next(sh for sh in s.shapes if sh.name == "Rectangle 4")
paras = info.text_frame.paragraphs
texts = [pg.text for pg in paras]
for pg, t in zip(paras, [texts[1], texts[0], texts[2]]):
    pg.runs[0].text = t
    for r in pg.runs[1:]: r.text = ""
sub = s.shapes.add_textbox(Inches(0.9), Inches(3.38), Inches(8.2), Inches(0.5))
pg = sub.text_frame.paragraphs[0]; pg.alignment = PP_ALIGN.CENTER
r = pg.add_run(); r.text = "- 측정 주기와 지연을 고려한 카메라·IMU·거리센서 융합 추정기 -"; font(r, 18, False, NAVY)
for sh in s.shapes:
    if sh.has_text_frame:
        for pg in sh.text_frame.paragraphs:
            for r in pg.runs:
                rPr = r._r.get_or_add_rPr()
                for e in rPr.findall(qn("a:ea")): rPr.remove(e)
                etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)
notes(s, "[10초]\n안녕하세요, 학생자율연구II 연구계획을 발표할 박찬혁입니다. 제 연구는 GPS를 쓰기 어려운 환경에서 카메라, IMU, 거리센서를 하나의 필터로 융합하되, 센서마다 다른 측정 시점까지 고려하는 위치·자세 추정기를 만드는 것입니다.")

# ------------------------------------------------------------------ 3 연구 배경
s = S[2]; clear(s); title(s, "연구 배경")
body(s, 0.5, 1.12, 9.0, 5.6, [
    ("sub", "차량과 항공기의 항법: 예측과 보정"),
    ("sec", "항법의 기본 구조"),
    ("b", "GPS를 사용할 수 없는 터널·실내에서는 **IMU로 움직임을 예측**하고 **외부 센서로 보정**하여 위치를 추정함"),
    ("b", "차량 내비게이션은 터널에서 GPS가 끊겨도 속도·회전 정보로 위치를 이어서 계산함"),
    ("sec", "카메라를 이용한 보정: Optical Flow와 VIO"),
    ("b", "**Optical Flow**는 연속 영상에서 점의 이동을 추적하여 **영상 속 움직임을 측정**함 (Lucas and Kanade, 1981)"),
    ("b", "**VIO**는 영상 측정을 IMU와 결합하여 **위치·자세를 추정**함 (Mourikis and Roumeliotis, 2007; Geneva et al., 2020)"),
    ("sec", "차량에서 드론·항공기로"),
    ("b", "실내 드론은 하향 카메라·IMU·거리센서로 GPS 없이 위치를 유지함"),
    ("b", "화성 헬리콥터 Ingenuity도 **하향 카메라·IMU·레이저 고도계**를 EKF로 융합하여 비행함 (Bayard et al., 2019)"),
    ("b", "플랫폼은 달라도 **IMU로 예측하고 카메라·거리센서로 보정하는 구조**는 동일함"),
])
notes(s, "[65초]\n자동차 내비게이션은 터널에 들어가 GPS가 끊겨도 화면 속 차량이 계속 움직입니다. 마지막 위치에서 속도와 회전 정보로 위치를 이어서 계산하기 때문입니다. 이처럼 항법은 관성 정보로 움직임을 예측하고, 다른 센서로 그 예측을 보정하는 구조입니다.\n"
         "보정에 많이 쓰이는 센서가 카메라입니다. 옵티컬 플로우는 연속된 영상에서 점이 얼마나 움직였는지 추적하는 기법으로, 광마우스가 움직임을 읽는 것과 같은 원리입니다. VIO는 이 영상 정보를 IMU와 결합해 위치와 자세를 추정합니다. 옵티컬 플로우가 측정이라면, VIO는 그 측정을 쓰는 추정입니다.\n"
         "이 원리는 드론과 항공기에도 그대로 쓰입니다. 실내 드론은 하향 카메라와 IMU, 거리센서로 위치를 유지하고, NASA의 화성 헬리콥터 인저뉴어티도 같은 조합으로 비행했습니다. 플랫폼은 달라도 IMU로 예측하고, 카메라와 거리센서로 보정한다는 구조는 같습니다.")

# ------------------------------------------------------------------ 4 연구의 필요성 및 목적
s = S[3]; clear(s); title(s, "연구의 필요성 및 목적")
body(s, 0.5, 1.12, 9.0, 4.55, [
    ("sub", "실제 센서 융합에서 생기는 두 가지 문제"),
    ("sec", "개별 센서의 한계"),
    ("b", "IMU만 사용하면 적분 과정에서 **작은 오차가 시간에 따라 누적**됨"),
    ("b", "카메라는 영상 속 이동만 측정하므로 **실제 이동 거리로 환산하려면 거리 정보가 필요**함"),
    ("b", "따라서 IMU로 예측하고 카메라·거리센서로 보정하는 **센서 융합이 필요**함"),
    ("sec", "센서별 측정 시점의 차이"),
    ("b", "IMU는 빠른 주기, 카메라는 느린 주기로 측정되며 **영상 처리에도 시간이 소요**됨"),
    ("b", "필터에 도착한 영상 측정은 **도착 시점보다 이전의 움직임**을 나타냄"),
    ("b", "e.g. 5 m/s로 이동 중 50 ms 늦은 측정을 현재 값으로 쓰면 **약 25 cm 차이**가 발생함 (단순 계산 예시)"),
])
goal = s.shapes.add_shape(1, Inches(0.5), Inches(5.05), Inches(9.0), Inches(0.98))
goal.fill.solid(); goal.fill.fore_color.rgb = RGBColor(0xEE, 0xF0, 0xF6)
goal.line.color.rgb = NAVY; goal.line.width = Pt(1); goal.shadow.inherit = False
tf = goal.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for side in ("left", "right"): setattr(tf, "margin_" + side, Inches(0.2))
pg = tf.paragraphs[0]; pg.alignment = PP_ALIGN.LEFT
r = pg.add_run(); r.text = "연구 목적  "; font(r, 16, True, NAVY)
for t, b in runs_of("측정 주기와 지연 차이를 고려하여 카메라·IMU·거리센서를 **하나의 필터로 융합하는 위치·자세 추정기**를 구현함"):
    r = pg.add_run(); r.text = t; font(r, 15, b, BLACK if b else TEXT)
notes(s, "[55초]\n그런데 이 구조를 실제 센서에 적용하면 두 가지 문제가 생깁니다.\n"
         "첫째는 개별 센서의 한계입니다. IMU만 쓰면 작은 오차가 적분되면서 계속 쌓입니다. 카메라는 영상 속 이동만 알려 주기 때문에, 실제 이동 거리로 바꾸려면 거리센서 정보가 필요합니다. 그래서 세 센서를 함께 써야 합니다.\n"
         "둘째는 측정 시점의 차이입니다. IMU는 매우 빠른 주기로 들어오지만 카메라는 그보다 느리고, 영상 처리에도 시간이 걸립니다. 그래서 지금 필터에 들어온 영상 측정은 사실 조금 전의 움직임입니다. 예를 들어 초속 5 m로 움직일 때 50 ms 늦은 측정을 현재 값처럼 쓰면 25 cm의 차이가 생깁니다.\n"
         "그래서 제 연구의 목적은 측정 주기와 지연 차이를 고려해 카메라, IMU, 거리센서를 하나의 필터로 융합하는 위치·자세 추정기를 구현하는 것입니다.")

# ------------------------------------------------------------------ 5 선행연구와 적용 방향
s = S[4]; clear(s); title(s, "선행연구와 적용 방향")
body(s, 0.5, 1.12, 9.0, 5.6, [
    ("sub", "차량 분야의 비동기 센서 처리 방법을 비전 항법에 적용"),
    ("sec", "레이더·V2X 기반 선행 차량 상태 추정 (Kim and Ahn, 2025)"),
    ("b", "도착 시점이 서로 다른 레이더와 V2X 정보를 융합하여 앞 차량의 상태를 추정함"),
    ("b", "본 연구에서는 다음 세 가지 **시간 처리 방법**을 참고함"),
    ("b2", "**지연 보상**: 데이터의 수신 시점과 실제 측정 시점을 구분"),
    ("b2", "**시간 정렬**: 함께 사용하는 신호의 시점을 맞추고, 필요한 경우 보간"),
    ("b2", "**과거 상태 재추정**: 지연된 측정으로 과거 추정값을 보정한 후 현재까지 다시 추정"),
    ("sec", "비전 항법 분야의 참고 구조"),
    ("b", "MSCKF와 OpenVINS는 **과거 카메라 자세를 상태에 유지**하며 추정함 (Mourikis and Roumeliotis, 2007; Geneva et al., 2020)"),
    ("sec", "본 연구의 적용 방향"),
    ("b", "차량 모델을 그대로 쓰지 않고, **시간 처리 방법을 카메라·IMU·거리센서 구성에 맞게 적용**함"),
    ("b", "자체 센서 구성의 추정기를 구현하고 **시간 처리 적용 전·후의 차이**를 확인함"),
])
notes(s, "[55초]\n이 시간 문제는 차량 분야에서도 중요하게 다뤄졌습니다. 지도교수님 연구실의 Kim과 Ahn의 논문은 레이더와 V2X 통신처럼 도착 시점이 서로 다른 정보를 융합해 앞 차량의 상태를 추정했습니다. 저는 이 논문에서 세 가지 처리 방법을 참고하겠습니다.\n"
         "첫째, 지연 보상입니다. 데이터를 받은 시점과 실제로 측정한 시점을 구분합니다. 둘째, 시간 정렬입니다. 함께 쓰는 신호의 시점을 맞추고, 필요하면 앞뒤 데이터로 그 사이 값을 보간합니다. 셋째, 과거 상태 재추정입니다. 늦게 들어온 측정으로 그 시점의 추정값을 고치고, 거기서부터 현재까지 다시 계산합니다.\n"
         "비전 항법 쪽에서는 과거 카메라 자세를 함께 유지하며 추정하는 MSCKF와 OpenVINS의 구조를 참고하겠습니다. 차량 모델을 그대로 가져오는 것이 아니라, 시간 처리 방법을 제 센서 구성에 맞게 옮기는 것이 핵심입니다.")

# ------------------------------------------------------------------ 6 연구방법
s = S[5]; clear(s); title(s, "연구방법")
body(s, 0.5, 1.12, 9.0, 5.6, [
    ("sub", "EKF 기반 다중센서 융합과 지연 측정 처리"),
    ("sec", "기본 융합 구조"),
    ("b", "영상 이동 추적은 **KLT Optical Flow**, 센서 융합은 **EKF(확장 칼만필터)**를 기본 구조로 사용함"),
    ("b", "**예측**: IMU의 가속도·각속도로 다음 시점의 위치·자세를 계산"),
    ("b", "**보정**: 영상 이동과 거리 정보를 측정 모델로 반영 (자이로로 회전 성분 제거, 거리로 실제 크기 결정)"),
    ("sec", "시간 처리"),
    ("b", "센서별 **측정 주기·타임스탬프·지연**을 먼저 확인함"),
    ("b", "과거 일정 구간의 **상태·IMU 입력·측정값을 저장**하고, 새 측정이 없으면 예측만 진행함"),
    ("b", "늦게 도착한 측정은 **측정 시점으로 돌아가 보정한 후, 저장한 IMU 입력으로 현재까지 재추정**함"),
    ("sec", "검증 계획"),
    ("b", "같은 데이터·초기 조건에서 **시간 처리 적용 전·후를 비교**함"),
    ("b", "기준값 대비 **위치·자세 오차**와 재추정에 필요한 **처리시간**을 함께 분석함"),
])
notes(s, "[65초]\n구체적인 방법입니다. 영상의 이동은 KLT 옵티컬 플로우로 추적하고, 융합은 확장 칼만필터, 즉 EKF를 기본 구조로 사용합니다.\n"
         "예측 단계에서는 IMU의 가속도와 각속도로 다음 순간의 위치와 자세를 계산합니다. 보정 단계에서는 영상 이동과 거리 정보를 측정 모델로 넣어 예측을 바로잡습니다. 이때 영상 이동에는 회전 성분이 섞여 있어서, 자이로로 회전을 빼고 거리로 실제 크기를 정합니다.\n"
         "시간 처리는 이렇게 합니다. 먼저 센서별 주기와 타임스탬프, 지연을 확인합니다. 필터는 과거 일정 구간의 상태와 입력을 저장해 두고, 새 측정이 없을 때는 예측만 합니다. 늦게 도착한 측정이 오면 그 측정 시점으로 돌아가 보정하고, 저장해 둔 IMU 입력으로 현재까지 다시 계산합니다.\n"
         "검증은 같은 데이터와 초기 조건에서 시간 처리를 넣기 전과 후를 비교합니다. 기준값 대비 위치 오차와 재추정에 드는 처리시간을 함께 보고, 효과와 한계를 정리하겠습니다.")

# ------------------------------------------------------------------ 7 연구계획 (clean table)
s = S[6]
old = next(sh for sh in s.shapes if sh.has_table)
plan = [[old.table.cell(i, j).text_frame.text.replace("\x0b", "\n") for j in range(3)] for i in range(1, 5)]
clear(s); title(s, "연구계획")
body(s, 0.5, 1.12, 9.0, 0.45, [("sub", "10월 ~ 12월 단계별 추진 계획")])
rows, cols = 5, 3
gt = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.7), Inches(9.0), Inches(3.9))
tbl = gt.table
tblPr = tbl._tbl.tblPr
for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
sid = tblPr.find(qn("a:tableStyleId"))
if sid is not None: tblPr.remove(sid)
for c, w in zip(tbl.columns, [1.3, 4.4, 3.3]): c.width = Inches(w)
tbl.rows[0].height = Inches(0.45)
for i in range(1, 5): tbl.rows[i].height = Inches(0.86)


def border(tc, side, w, color):
    tcPr = tc._tc.get_or_add_tcPr()
    ln = etree.SubElement(tcPr, qn("a:" + side)); ln.set("w", str(int(Pt(w))))
    if w == 0:
        etree.SubElement(ln, qn("a:noFill")); return
    sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", color)


def fill_cell(cell, paras, header=False):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
    tf = cell.text_frame; tf.word_wrap = True
    for k, (t, b) in enumerate(paras):
        pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pg.line_spacing = 1.1
        r = pg.add_run(); r.text = t
        font(r, 14 if header else 13.5, b, NAVY if header else (BLACK if b else TEXT))
    tcPr = cell._tc.get_or_add_tcPr()
    for side, w, col in (("lnL", 0, None), ("lnR", 0, None),
                         ("lnT", 1.5 if header else 0.5, "000066" if header else "BFBFBF"),
                         ("lnB", 1.0 if header else 0.5, "000066" if header else "BFBFBF")):
        border(cell, side, w, col)
    if header:
        sf = etree.SubElement(tcPr, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", "E8ECF5")
    else:
        etree.SubElement(tcPr, qn("a:noFill"))


for j, h in enumerate(["시기", "수행 내용", "확인할 결과"]):
    fill_cell(tbl.cell(0, j), [(h, True)], header=True)
for i, (when, what, res) in enumerate(plan, start=1):
    fill_cell(tbl.cell(i, 0), [(when.replace("\n", " "), True)])
    wl = [l for l in what.split("\n") if l.strip()]
    fill_cell(tbl.cell(i, 1), [(wl[0], True)] + [(l, False) for l in wl[1:]])
    rl = [l for l in res.split("\n") if l.strip()]
    fill_cell(tbl.cell(i, 2), [(l, False) for l in rl])
body(s, 0.5, 5.85, 9.0, 0.8, [
    ("sec", "최종 목표"),
    ("b", "**시간 처리를 포함한 다중센서 융합 추정기 구현**과 시간 처리 효과 검증"),
])
notes(s, "[25초]\n일정입니다. 10월 전반에는 선행논문을 분석하고 센서의 주기와 타임스탬프, 지연을 정리합니다. 10월 후반에는 옵티컬 플로우와 IMU 예측으로 기본 융합 구조를 만듭니다. 11월에는 지연 보상과 시간 정렬, 과거 상태 재추정을 넣고, 12월에 적용 전후를 비교해 보고서로 정리하겠습니다.")

# ------------------------------------------------------------------ 8 참고문헌 / 감사합니다
s = S[7]; clear(s); title(s, "참고문헌 / 감사합니다")
refs = [
    "[1] M. Kim and C. Ahn, “State Estimation of Preceding Target Vehicle Using Radar and V2X,” IEEE Access, vol. 13, pp. 198482–198495, 2025.",
    "[2] A. I. Mourikis and S. I. Roumeliotis, “A Multi-State Constraint Kalman Filter for Vision-aided Inertial Navigation,” IEEE ICRA, pp. 3565–3572, 2007.",
    "[3] P. Geneva et al., “OpenVINS: A Research Platform for Visual-Inertial Estimation,” IEEE ICRA, pp. 4666–4672, 2020.",
    "[4] D. S. Bayard et al., “Vision-Based Navigation for the NASA Mars Helicopter,” AIAA SciTech Forum, 2019.",
    "[5] B. D. Lucas and T. Kanade, “An Iterative Image Registration Technique with an Application to Stereo Vision,” IJCAI, 1981.",
]
tb = s.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(9.0), Inches(3.6))
tf = tb.text_frame; tf.word_wrap = True
for k, t in enumerate(refs):
    pg = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    bullet(pg, None, 0.35, 0.35); spacing(pg, 0 if k == 0 else 6, 0, 1.1)
    r = pg.add_run(); r.text = t; font(r, 13.5, False, TEXT)
body(s, 0.5, 4.95, 9.0, 1.7, [
    ("sub", "감사합니다"),
    ("p", "연구계획에 대한 의견과 질문 부탁드립니다."),
])
notes(s, "[15초]\n정리하면, 차량 분야에서 다뤄 온 시간 처리 방법을 카메라·IMU·거리센서 기반 비전 항법에 적용하고 그 효과를 확인하는 것이 이 연구입니다. 이상으로 발표를 마치겠습니다. 의견과 질문 부탁드립니다. 감사합니다.")

# ------------------------------------------------------------------ delete 목차
lst = p.slides._sldIdLst
sid = lst.findall(qn("p:sldId"))[1]
p.part.drop_rel(sid.get(qn("r:id"))); lst.remove(sid)
p.save(OUT)
print("saved", OUT)
